"""Generate the synthetic sample decks used by the tests and the README.

Run with:  python sample_data/generate_samples.py

Three decks, chosen to exercise genuinely different paths through the model:

* northwind_saas_series_a.pptx - revenue-mode sector, strong round progress
* helix_bio_preseed.pdf     - milestone-mode sector, pre-revenue by design
* sparse_deck.pdf           - almost no disclosure, to prove the model degrades
                              gracefully instead of inventing values
"""

from __future__ import annotations

import sys
from pathlib import Path

SAMPLE_DIR = Path(__file__).resolve().parent

# --------------------------------------------------------------------------
# Deck content
# --------------------------------------------------------------------------
NORTHWIND_SLIDES: list[tuple[str, list[str]]] = [
    ("Northwind Analytics", ["Revenue intelligence for industrial distributors", "www.northwindanalytics.com"]),
    ("The Problem", [
        "Industrial distributors lose 12% of gross margin to unmanaged pricing.",
        "The average distributor runs pricing in 14 disconnected spreadsheets.",
    ]),
    ("Our Solution", [
        "A SaaS pricing platform that plugs into existing ERP systems.",
        "Deployed in 6 weeks, no data warehouse required.",
        "Business model: per-seat subscription software with annual contracts.",
    ]),
    ("Traction", [
        "ARR: $2,400,000",
        "MRR: $200,000",
        "Revenue growth: 145% year-over-year",
        "42 paying customers, including 11 enterprise customers",
        "Net revenue retention: 118%",
        "Gross margin: 82%",
        "Average contract value: $57,000",
    ]),
    ("Market", [
        "TAM: $8.4 billion of industrial distribution software spend",
        "SAM: $1.9 billion in North America",
        "SOM: $190 million within five years",
        "Market growth: 14% CAGR",
    ]),
    ("Team", [
        "Dana Whitfield - Co-Founder & CEO",
        "Serial entrepreneur; previously founded Lumen Pricing (acquired by SAP in 2019).",
        "Commitment: Full-Time",
        "Marcus Ilori - Co-Founder & CTO",
        "15 years of pricing systems engineering.",
        "Priya Raman - Chief Revenue Officer",
        "Scaled enterprise sales from $2M to $40M at a prior company.",
        "38 employees",
    ]),
    ("Defensibility", [
        "Two granted patents on dynamic margin optimisation (USPTO).",
        "Proprietary benchmark dataset across 340 distributors.",
        "Competitive landscape: legacy pricing suites with 9-month implementations.",
    ]),
    ("Financials", [
        "Monthly burn rate: $310,000",
        "Cash balance: $5,600,000",
        "Runway: 18 months",
        "Financial model and three-year projections available in the data room.",
    ]),
    ("The Round", [
        "Currently raising $9,000,000 Series A",
        "Pre-money valuation: $36,000,000",
        "Security type: Series A Preferred Stock",
        "$5,400,000 committed - 60% of the round is committed",
        "Lead investor: Calder Ridge Ventures",
        "Existing investors: Foundry Park Capital, Tessellate Ventures",
        "Prior funding raised: $4,200,000",
        "Minimum investment: $250,000",
    ]),
    ("Use of Funds", [
        "Use of proceeds: 45% engineering, 35% go-to-market, 12% customer success, 8% G&A.",
        "Milestones: $6M ARR and 90 customers within 18 months.",
    ]),
]

HELIX_PAGES: list[tuple[str, list[str]]] = [
    ("Helix Bioworks", ["A first-in-class oral therapeutic for refractory autoimmune disease"]),
    ("Problem", [
        "2.4 million patients fail first-line biologic therapy each year.",
        "Annual cost to payors exceeds $14 billion.",
    ]),
    ("Science", [
        "Small molecule compound HB-114 targets a novel intracellular pathway.",
        "Preclinical efficacy demonstrated in two validated animal models.",
        "Technology readiness: TRL 4",
    ]),
    ("Regulatory", [
        "Pre-clinical stage; IND-enabling studies underway.",
        "Regulatory pathway: FDA 505(b)(1) with orphan drug designation sought.",
    ]),
    ("Intellectual Property", [
        "Composition-of-matter patent filed (PCT), exclusive licence from Northgate University.",
    ]),
    ("Team", [
        "Dr. Ines Moreau - Co-Founder & CEO",
        "Commitment: Full-Time",
        "Dr. Samuel Oyelaran - Co-Founder & Chief Scientific Officer",
        "9 employees",
    ]),
    ("Partnerships", [
        "Signed a joint development agreement with Marrow Clinical Research.",
    ]),
    ("Market", [
        "TAM: $22 billion refractory autoimmune market",
        "SAM: $4.1 billion addressable with an oral therapy",
    ]),
    ("The Round", [
        "Raising $3,500,000 pre-seed",
        "Security type: SAFE note",
        "Valuation cap: $12,000,000",
        "Discount: 20%",
        "Non-dilutive funding: $1,100,000 NIH SBIR award",
        "Prior funding raised: $750,000",
        "Runway: 14 months",
        "Monthly burn rate: $180,000",
        "Use of proceeds: IND-enabling toxicology, CMC and a Phase 1 submission.",
    ]),
]

SPARSE_PAGES: list[tuple[str, list[str]]] = [
    ("Orbital Freight", ["Logistics, reimagined."]),
    ("Why now", ["The market is changing fast."]),
    ("Contact", ["hello@orbitalfreight.io"]),
]


# --------------------------------------------------------------------------
# Writers
# --------------------------------------------------------------------------
def write_pptx(path: Path, slides: list[tuple[str, list[str]]]) -> Path:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    presentation = Presentation()
    blank = presentation.slide_layouts[6]
    for title, lines in slides:
        slide = presentation.slides.add_slide(blank)

        title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(8.8), Inches(0.9))
        frame = title_box.text_frame
        frame.text = title
        frame.paragraphs[0].runs[0].font.size = Pt(30)
        frame.paragraphs[0].runs[0].font.bold = True

        body = slide.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(8.8), Inches(4.8))
        body_frame = body.text_frame
        body_frame.word_wrap = True
        for index, line in enumerate(lines):
            paragraph = body_frame.paragraphs[0] if index == 0 else body_frame.add_paragraph()
            paragraph.text = line
            paragraph.runs[0].font.size = Pt(16)

        # Speaker notes exercise the notes-extraction path.
        slide.notes_slide.notes_text_frame.text = f"Speaker notes for {title}."

    path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(path))
    return path


def write_pdf(path: Path, pages: list[tuple[str, list[str]]], letter_spaced_title: bool = False) -> Path:
    from reportlab.lib.pagesizes import LETTER
    from reportlab.pdfgen import canvas as pdfcanvas

    path.parent.mkdir(parents=True, exist_ok=True)
    width, height = LETTER
    canvas = pdfcanvas.Canvas(str(path), pagesize=LETTER)
    for title, lines in pages:
        canvas.setFont("Helvetica-Bold", 24)
        # One deck deliberately uses tracked-out headings, which is how real
        # design decks defeat naive keyword extraction.
        canvas.drawString(56, height - 80, " ".join(title) if letter_spaced_title else title)
        canvas.setFont("Helvetica", 13)
        y = height - 130
        for line in lines:
            canvas.drawString(56, y, line)
            y -= 22
        canvas.showPage()
    canvas.save()
    return path


def main() -> int:
    written = [
        write_pptx(SAMPLE_DIR / "northwind_saas_series_a.pptx", NORTHWIND_SLIDES),
        write_pdf(SAMPLE_DIR / "helix_bio_preseed.pdf", HELIX_PAGES, letter_spaced_title=True),
        write_pdf(SAMPLE_DIR / "sparse_deck.pdf", SPARSE_PAGES),
    ]
    unsupported = SAMPLE_DIR / "unsupported.txt"
    unsupported.write_text("This is not a pitch deck.\n", encoding="utf-8")
    written.append(unsupported)

    for path in written:
        print(f"wrote {path.relative_to(SAMPLE_DIR.parent)}  ({path.stat().st_size:,} bytes)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
