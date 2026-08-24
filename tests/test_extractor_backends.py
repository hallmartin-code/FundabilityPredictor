"""Format-backend edge cases: fallbacks, corrupt files and legacy .ppt."""

from __future__ import annotations

import pytest

from extractors import ppt_extractor
from extractors.base import ExtractionError, Page
from extractors.ingest import extractor_for, ingest
from extractors.pdf_extractor import PdfExtractor
from extractors.ppt_extractor import LegacyPptExtractor, PptxExtractor


# ---------------------------------------------------------------------------
# Dispatch
# ---------------------------------------------------------------------------
@pytest.mark.parametrize(
    ("filename", "expected"),
    [
        ("deck.pdf", PdfExtractor),
        ("deck.PDF", PdfExtractor),
        ("deck.pptx", PptxExtractor),
        ("deck.pptm", PptxExtractor),
        ("deck.ppt", LegacyPptExtractor),
    ],
)
def test_extension_dispatch(filename, expected, tmp_path):
    assert isinstance(extractor_for(tmp_path / filename), expected)


# ---------------------------------------------------------------------------
# PDF backend
# ---------------------------------------------------------------------------
def test_pdf_falls_back_to_pypdf_when_pdfplumber_fails(biotech_deck, monkeypatch):
    calls: list[str] = []

    def _fail(self, path):
        calls.append("pdfplumber")
        return None

    monkeypatch.setattr(PdfExtractor, "_parse_with_pdfplumber", _fail)
    document = PdfExtractor().extract(biotech_deck)

    assert calls == ["pdfplumber"]
    assert document.pages
    # Body text is unspaced in this sample; the headings are deliberately not.
    assert "Monthly burn rate: $180,000" in document.full_text


def test_corrupt_pdf_raises_extraction_error(tmp_path):
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"%PDF-1.4\nnot actually a pdf at all\n")
    with pytest.raises(ExtractionError):
        ingest(broken)


def test_pdf_page_failures_do_not_lose_the_document(biotech_deck, monkeypatch):
    """One unreadable page must not discard the rest of the deck."""
    import pdfplumber

    original = pdfplumber.page.Page.extract_text
    state = {"calls": 0}

    def _flaky(self, *args, **kwargs):
        state["calls"] += 1
        if state["calls"] == 2:
            raise RuntimeError("simulated page failure")
        return original(self, *args, **kwargs)

    monkeypatch.setattr(pdfplumber.page.Page, "extract_text", _flaky)
    document = PdfExtractor().extract(biotech_deck)
    assert len(document.pages) >= 8
    assert document.full_text.strip()


# ---------------------------------------------------------------------------
# PowerPoint backend
# ---------------------------------------------------------------------------
def test_pptx_tables_and_notes_are_extracted(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(6), Inches(1.4))
    table = shape.table
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "ARR"
    table.cell(1, 1).text = "$2,400,000"
    slide.notes_slide.notes_text_frame.text = "Runway: 18 months"

    path = tmp_path / "tabular.pptx"
    presentation.save(str(path))

    document = ingest(path)
    page = document.pages[0]
    assert page.tables
    assert "ARR | $2,400,000" in page.table_text()
    # Tables and notes both feed the searchable text used by the extractors.
    assert "$2,400,000" in page.searchable_text()
    assert "Runway: 18 months" in page.searchable_text()


def test_pptx_charts_contribute_their_labels(tmp_path):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["2024", "2025", "2026"]
    data.add_series("ARR", (1.0, 2.4, 5.0))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(6), Inches(4), data
    )

    path = tmp_path / "charted.pptx"
    presentation.save(str(path))

    text = ingest(path).pages[0].searchable_text()
    assert "2024" in text
    assert "ARR" in text


def test_empty_presentation_raises(tmp_path):
    from pptx import Presentation

    path = tmp_path / "empty.pptx"
    Presentation().save(str(path))
    with pytest.raises(ExtractionError):
        ingest(path)


def test_legacy_ppt_without_libreoffice_gives_actionable_guidance(tmp_path, monkeypatch):
    monkeypatch.setattr(ppt_extractor, "_find_soffice", lambda: None)
    legacy = tmp_path / "old_deck.ppt"
    legacy.write_bytes(b"\xd0\xcf\x11\xe0legacy binary powerpoint")

    with pytest.raises(ExtractionError) as excinfo:
        ingest(legacy)

    message = str(excinfo.value)
    assert "LibreOffice" in message
    assert ".pptx" in message and ".pdf" in message


def test_legacy_ppt_uses_libreoffice_when_available(tmp_path, monkeypatch, saas_deck):
    """The conversion step hands off to the .pptx reader."""
    converted = tmp_path / "converted" / "old_deck.pptx"
    converted.parent.mkdir(parents=True)
    converted.write_bytes(saas_deck.read_bytes())

    monkeypatch.setattr(ppt_extractor, "_convert_to_pptx", lambda path: converted)
    legacy = tmp_path / "old_deck.ppt"
    legacy.write_bytes(b"\xd0\xcf\x11\xe0legacy binary powerpoint")

    document = ingest(legacy)
    assert "Northwind" in document.full_text


# ---------------------------------------------------------------------------
# Quality reporting
# ---------------------------------------------------------------------------
def test_image_based_pages_are_flagged():
    from extractors.base import ExtractedDocument
    from models.profile import SourceDocument

    document = ExtractedDocument(
        source=SourceDocument(path="x", filename="x.pdf", file_type="pdf"),
        pages=[
            Page(number=1, text="A" * 400),
            Page(number=2, text=""),
            Page(number=3, text="short"),
        ],
    )
    quality = document.quality()
    assert quality.image_based_pages == [2, 3]
    assert quality.pages_with_text == 1
    assert any("IMAGE-BASED CONTENT" in w for w in quality.warnings())


def test_image_heavy_documents_are_identified():
    from extractors.base import ExtractedDocument
    from models.profile import SourceDocument

    document = ExtractedDocument(
        source=SourceDocument(path="x", filename="x.pdf", file_type="pdf"),
        pages=[Page(number=n, text="" if n > 2 else "A" * 400) for n in range(1, 6)],
    )
    assert document.quality().is_image_heavy


def test_source_hash_is_recorded(saas_deck):
    document = ingest(saas_deck)
    assert document.source.sha256
    assert len(document.source.sha256) == 64
    # Re-reading the same file must produce the same digest.
    assert ingest(saas_deck).source.sha256 == document.source.sha256
