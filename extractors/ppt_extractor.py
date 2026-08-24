"""PowerPoint ingestion.

``.pptx`` is read natively with python-pptx. Legacy binary ``.ppt`` is not a
readable format for python-pptx, so it is converted to ``.pptx`` first using a
locally installed LibreOffice; if none is available the user gets an
actionable error rather than a silent empty extraction.
"""

from __future__ import annotations

import logging
import shutil
import subprocess
import tempfile
from pathlib import Path

from extractors.base import DocumentExtractor, ExtractionError, Page

LOGGER = logging.getLogger(__name__)

#: Executables that can convert legacy .ppt to .pptx, in preference order.
_SOFFICE_CANDIDATES = (
    "soffice",
    "libreoffice",
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    "/usr/bin/soffice",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
)

_CONVERSION_TIMEOUT_SECONDS = 180


class PptxExtractor(DocumentExtractor):
    extensions = (".pptx", ".pptm")
    file_type = "pptx"

    def parse(self, path: Path) -> list[Page]:
        try:
            from pptx import Presentation
        except ImportError as exc:  # pragma: no cover - dependency is declared
            raise ExtractionError(
                "python-pptx is not installed; cannot read PowerPoint files."
            ) from exc

        try:
            presentation = Presentation(str(path))
        except Exception as exc:  # noqa: BLE001
            raise ExtractionError(f"Could not open {path.name}: {exc}") from exc

        pages: list[Page] = []
        for index, slide in enumerate(presentation.slides, start=1):
            texts: list[str] = []
            tables: list[list[list[str]]] = []
            for shape in slide.shapes:
                texts.extend(_shape_text(shape))
                table = _shape_table(shape)
                if table:
                    tables.append(table)
            pages.append(
                Page(
                    number=index,
                    text="\n".join(t for t in texts if t.strip()),
                    tables=tables,
                    notes=_slide_notes(slide),
                    title=_slide_title(slide),
                )
            )
        if not pages:
            raise ExtractionError(f"{path.name} contains no slides.")
        return pages


class LegacyPptExtractor(PptxExtractor):
    """Legacy ``.ppt`` support by way of a LibreOffice conversion step."""

    extensions = (".ppt",)
    file_type = "ppt"

    def parse(self, path: Path) -> list[Page]:
        converted = _convert_to_pptx(path)
        try:
            return super().parse(converted)
        finally:
            shutil.rmtree(converted.parent, ignore_errors=True)


def _convert_to_pptx(path: Path) -> Path:
    soffice = _find_soffice()
    if not soffice:
        raise ExtractionError(
            f"{path.name} is a legacy binary .ppt file, which cannot be read directly. "
            "Install LibreOffice (so that 'soffice' is on PATH) to enable automatic "
            "conversion, or re-save the deck as .pptx or .pdf and try again."
        )
    tmpdir = Path(tempfile.mkdtemp(prefix="fp_ppt_"))
    try:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pptx", "--outdir", str(tmpdir), str(path)],
            check=True,
            capture_output=True,
            timeout=_CONVERSION_TIMEOUT_SECONDS,
        )
    except subprocess.TimeoutExpired as exc:
        shutil.rmtree(tmpdir, ignore_errors=True)
        raise ExtractionError(f"Converting {path.name} timed out.") from exc
    except (subprocess.CalledProcessError, OSError) as exc:
        shutil.rmtree(tmpdir, ignore_errors=True)
        raise ExtractionError(f"Could not convert {path.name} to .pptx: {exc}") from exc

    produced = next(iter(tmpdir.glob("*.pptx")), None)
    if produced is None:
        shutil.rmtree(tmpdir, ignore_errors=True)
        raise ExtractionError(f"Conversion of {path.name} produced no .pptx output.")
    LOGGER.info("Converted legacy .ppt %s via LibreOffice.", path.name)
    return produced


def _find_soffice() -> str | None:
    for candidate in _SOFFICE_CANDIDATES:
        resolved = shutil.which(candidate)
        if resolved:
            return resolved
        if Path(candidate).exists():
            return candidate
    return None


# ---------------------------------------------------------------------------
# Shape readers. Each is defensive: one malformed shape must not lose a slide.
# ---------------------------------------------------------------------------
def _shape_text(shape) -> list[str]:  # noqa: ANN001 - pptx shape is untyped
    out: list[str] = []
    try:
        if shape.shape_type is not None and getattr(shape, "has_text_frame", False):
            text = shape.text_frame.text
            if text and text.strip():
                out.append(text.strip())
    except Exception as exc:  # noqa: BLE001
        LOGGER.debug("Could not read shape text: %s", exc)

    try:
        if getattr(shape, "has_chart", False):
            out.extend(_chart_text(shape.chart))
    except Exception as exc:  # noqa: BLE001
        LOGGER.debug("Could not read chart: %s", exc)

    # Grouped shapes nest arbitrarily deep.
    try:
        if getattr(shape, "shape_type", None) is not None and hasattr(shape, "shapes"):
            for child in shape.shapes:
                out.extend(_shape_text(child))
    except Exception as exc:  # noqa: BLE001
        LOGGER.debug("Could not walk group shape: %s", exc)
    return out


def _chart_text(chart) -> list[str]:  # noqa: ANN001
    """Chart titles, category labels and series names, where available."""
    out: list[str] = []
    try:
        if chart.has_title and chart.chart_title.text_frame.text.strip():
            out.append(chart.chart_title.text_frame.text.strip())
    except Exception:  # noqa: BLE001
        pass
    try:
        categories = [str(c) for c in chart.plots[0].categories if c is not None]
        if categories:
            out.append(" | ".join(categories))
    except Exception:  # noqa: BLE001
        pass
    try:
        for series in chart.series:
            name = getattr(series, "name", None)
            values = [v for v in getattr(series, "values", []) or [] if v is not None]
            if name:
                rendered = ", ".join(f"{v:g}" for v in values[:12])
                out.append(f"{name}: {rendered}" if rendered else str(name))
    except Exception:  # noqa: BLE001
        pass
    return out


def _shape_table(shape) -> list[list[str]] | None:  # noqa: ANN001
    try:
        if not getattr(shape, "has_table", False):
            return None
        rows = []
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                rows.append(cells)
        return rows or None
    except Exception as exc:  # noqa: BLE001
        LOGGER.debug("Could not read table: %s", exc)
        return None


def _slide_notes(slide) -> str:  # noqa: ANN001
    try:
        if slide.has_notes_slide:
            return (slide.notes_slide.notes_text_frame.text or "").strip()
    except Exception as exc:  # noqa: BLE001
        LOGGER.debug("Could not read speaker notes: %s", exc)
    return ""


def _slide_title(slide) -> str:  # noqa: ANN001
    try:
        title = slide.shapes.title
        if title is not None and title.text_frame.text.strip():
            return title.text_frame.text.strip()[:120]
    except Exception:  # noqa: BLE001
        pass
    return ""
